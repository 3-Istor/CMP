import asyncio
import io
from datetime import date, datetime
from unittest import mock
from zoneinfo import ZoneInfo

import pyarrow.parquet as pq
import pytest
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.core.config import settings
from app.core.database import Base, get_db
from app.core.roles import ProjectRole
from app.models.finops import (
    CostCenter,
    CostLineItem,
    FinopsState,
    MonthClose,
    SavingsEntry,
)
from app.models.project import Project
from app.routers import finops as router
from app.routers import finops_finance
from app.services.audit import match_action
from app.services.authz import is_cnp_finance
from app.services.finops import finance, monthly_report
from app.services.finops.ingest.store import CostLine, upsert
from app.services.keycloak_service import get_current_user

ADMIN = {"preferred_username": "root"}
FINANCE = {"preferred_username": "fin", "groups": ["/cnp-finance"]}
MEMBER = {"preferred_username": "mem"}


@pytest.fixture
def factory():
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine)


@pytest.fixture(autouse=True)
def _roles(monkeypatch):
    monkeypatch.setattr(settings, "FINOPS_ADMIN_USERS", "root")
    monkeypatch.setattr(settings, "FINOPS_FINANCE_USERS", "")

    def only_qcm(_db, _token, project, _minimum):
        if project != "qcm":
            raise HTTPException(status_code=403)

    monkeypatch.setattr(router, "assert_project_role", only_qcm)
    monkeypatch.setattr(router, "_user_project_names", lambda token: ["qcm"])
    monkeypatch.setattr(
        finops_finance, "_user_project_names", lambda token: ["qcm"]
    )


@pytest.fixture
def api(factory):
    who = {"token": MEMBER}

    def _db():
        with factory() as db:
            yield db

    app = FastAPI()
    app.include_router(router.router, prefix="/api")
    app.include_router(finops_finance.router, prefix="/api")
    app.dependency_overrides[get_db] = _db
    app.dependency_overrides[get_current_user] = lambda: who["token"]
    client = TestClient(app)

    def as_user(token):
        who["token"] = token
        return client

    return as_user


def _line(
    project="qcm",
    cost=1.0,
    start=datetime(2026, 9, 10),
    resource="onprem/qcm/cpu",
    **extra,
):
    fields = dict(
        x_source="opencost",
        provider_name="CNP on-prem",
        service_name="Kubernetes CPU",
        service_category="Compute",
        charge_period_start=start,
        charge_period_end=start.replace(hour=23),
        resource_id=resource,
        x_cost_eur=cost,
        billed_cost=cost,
        list_cost=cost,
        effective_cost=cost,
        project=project,
        app="api",
        x_pricing_basis="tco",
        x_is_final=True,
    )
    fields.update(extra)
    return CostLine(**fields)


@pytest.fixture
def seeded(factory):
    with factory() as db:
        upsert(
            db,
            [
                _line("qcm", 2.0),
                _line("other", 3.0, resource="onprem/other/cpu"),
                _line(
                    "qcm",
                    5.0,
                    resource="=cmd|calc",
                    x_source="aws",
                    provider_name="AWS",
                    x_pricing_basis="invoice",
                    cloud="aws",
                    resource_name='=HYPERLINK("http://x")',
                ),
            ],
        )
        db.add_all(
            [
                Project(project_name="qcm", owner_username="mem"),
                Project(project_name="other", owner_username="x"),
                SavingsEntry(
                    day=date(2026, 9, 3),
                    project="qcm",
                    app="api",
                    category="offhours",
                    avoided_eur=4.0,
                    realized_eur=1.5,
                ),
                CostCenter(
                    project_name="qcm", cost_center="CC-1", budget_owner="Ana"
                ),
            ]
        )
        db.commit()
    return factory


# ── role detection ────────────────────────────────────────────────────────────


@pytest.mark.parametrize(
    "token",
    [
        {"groups": ["cnp-finance"]},
        {"groups": ["/cnp-finance"]},
        {"realm_access": {"roles": ["cnp-finance"]}},
        {"preferred_username": "Fin"},
    ],
)
def test_is_cnp_finance_accepts_group_role_or_listed_user(token, monkeypatch):
    monkeypatch.setattr(settings, "FINOPS_FINANCE_USERS", "fin, bob")

    assert is_cnp_finance(token) is True


def test_is_cnp_finance_rejects_other_groups():
    assert (
        is_cnp_finance({"groups": ["/cnp-finance-x", "/qcm-admins"]}) is False
    )


def test_meta_reports_the_callers_role(api):
    roles = {
        name: api(token).get("/api/finops/meta").json()["role"]
        for name, token in (
            ("admin", ADMIN),
            ("finance", FINANCE),
            ("member", MEMBER),
        )
    }

    assert roles == {
        "admin": "admin",
        "finance": "finance",
        "member": "member",
    }


# ── RBAC ──────────────────────────────────────────────────────────────────────


def _ledger(factory):
    with factory() as db:
        for project in ("qcm", "other"):
            db.add(
                SavingsEntry(
                    day=date.today(),
                    project=project,
                    app="a",
                    category="offhours",
                    avoided_eur=1.0,
                )
            )
        db.commit()


def test_finance_reads_savings_of_every_project(api, factory):
    _ledger(factory)

    body = api(FINANCE).get("/api/finops/savings").json()

    assert {r["project"] for r in body["rows"]} == {"qcm", "other"}


def test_finance_reads_a_project_it_is_not_member_of(api, factory):
    _ledger(factory)

    response = api(FINANCE).get("/api/finops/savings?project=other")

    assert response.status_code == 200


def test_member_savings_are_limited_to_their_projects(api, factory):
    _ledger(factory)

    body = api(MEMBER).get("/api/finops/savings").json()

    assert {r["project"] for r in body["rows"]} == {"qcm"}


def test_member_cannot_read_another_project(api, factory):
    assert (
        api(MEMBER).get("/api/finops/savings?project=other").status_code == 403
    )


def test_finance_cannot_set_a_budget(api):
    response = api(FINANCE).put(
        "/api/finops/budgets/qcm",
        json={
            "monthly_amount_eur": 10,
            "threshold_warn": 70,
            "threshold_critical": 90,
        },
    )

    assert response.status_code == 403


def test_finance_cannot_act_on_a_recommendation(factory):
    with factory() as db, pytest.raises(HTTPException) as err:
        router._require_project_role(db, FINANCE, "other", ProjectRole.ADMIN)

    assert err.value.status_code == 403


def test_admin_can_act_on_any_project(factory):
    with factory() as db:
        router._require_project_role(db, ADMIN, "other", ProjectRole.ADMIN)


@pytest.mark.parametrize(
    "method,url",
    [
        ("get", "/api/finops/export.xlsx?period=2026-09"),
        ("get", "/api/finops/export.parquet?period=2026-09"),
        ("post", "/api/finops/close/2026-09"),
    ],
)
def test_finance_only_endpoints_refuse_a_member(api, method, url):
    assert getattr(api(MEMBER), method)(url).status_code == 403


def test_member_cannot_write_cost_centers(api):
    response = api(MEMBER).put(
        "/api/finops/cost-centers",
        json={"centers": [{"project_name": "qcm", "cost_center": "X"}]},
    )

    assert response.status_code == 403


# ── exports ───────────────────────────────────────────────────────────────────


def test_xlsx_has_the_expected_sheets(api, seeded):
    response = api(FINANCE).get("/api/finops/export.xlsx?period=2026-09")

    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == [
        "Synthèse",
        "Projets",
        "Plateforme",
        "Fournisseurs",
        "Écarts",
        "Économies",
        "Lignes",
    ]


def test_xlsx_gaps_sheet_has_one_row_per_cloud(api, seeded):
    response = api(FINANCE).get("/api/finops/export.xlsx?period=2026-09")

    sheet = load_workbook(io.BytesIO(response.content))["Écarts"]
    assert [
        row[0] for row in sheet.iter_rows(min_row=3, values_only=True)
    ] == [
        "aws",
        "gcp",
    ]


def test_xlsx_covers_every_project_and_shows_the_pricing_basis(api, seeded):
    response = api(FINANCE).get("/api/finops/export.xlsx?period=2026-09")

    book = load_workbook(io.BytesIO(response.content))
    rows = [r for r in book["Synthèse"].iter_rows(values_only=True)]
    assert ("CNP on-prem", "tco", "OpEx", 5.0) in rows
    assert ("AWS", "invoice", "OpEx", 5.0) in rows
    assert any(r[0] == "other" for r in rows)


def test_xlsx_projects_sheet_includes_the_cost_center(api, seeded):
    response = api(FINANCE).get("/api/finops/export.xlsx?period=2026-09")

    sheet = load_workbook(io.BytesIO(response.content))["Projets"]
    assert ("qcm", "CC-1", "Ana") in [
        r[:3] for r in sheet.iter_rows(min_row=2, values_only=True)
    ]


def test_xlsx_never_sums_avoided_and_realized(api, seeded):
    response = api(FINANCE).get("/api/finops/export.xlsx?period=2026-09")

    sheet = load_workbook(io.BytesIO(response.content))["Économies"]
    total = next(
        r for r in sheet.iter_rows(values_only=True) if r[0] == "Total"
    )
    assert (total[4], total[5]) == (4.0, 1.5)


def test_xlsx_escapes_formulas(api, seeded):
    response = api(FINANCE).get("/api/finops/export.xlsx?period=2026-09")

    sheet = load_workbook(io.BytesIO(response.content))["Lignes"]
    cells = [
        c for row in sheet.iter_rows() for c in row if isinstance(c.value, str)
    ]
    assert not [c for c in cells if c.data_type == "f"]
    assert "'=cmd|calc" in {c.value for c in cells}


def test_parquet_has_the_focus_schema_and_the_month_lines(api, seeded):
    response = api(FINANCE).get("/api/finops/export.parquet?period=2026-09")

    table = pq.read_table(io.BytesIO(response.content))
    assert table.column_names == [
        name for name, _a, _t in finance.FOCUS_COLUMNS
    ]
    assert table.num_rows == 3
    assert str(table.schema.field("x_CostEur").type) == "double"


def test_parquet_excludes_other_months(api, seeded, factory):
    with factory() as db:
        upsert(db, [_line(start=datetime(2026, 8, 2))])

    response = api(FINANCE).get("/api/finops/export.parquet?period=2026-09")

    assert pq.read_table(io.BytesIO(response.content)).num_rows == 3


def test_export_refuses_a_bad_period(api):
    assert (
        api(FINANCE).get("/api/finops/export.xlsx?period=2026-13").status_code
        == 422
    )


def test_statement_pdf_is_a_non_empty_pdf_for_a_member(api, seeded):
    response = api(MEMBER).get(
        "/api/finops/statement.pdf?project=qcm&period=2026-09"
    )

    assert response.headers["content-type"] == "application/pdf"
    assert response.content.startswith(b"%PDF") and len(response.content) > 500


def test_statement_pdf_is_denied_on_a_foreign_project(api, seeded):
    response = api(MEMBER).get(
        "/api/finops/statement.pdf?project=other&period=2026-09"
    )

    assert response.status_code == 403


def test_statement_pdf_is_available_to_finance_for_any_project(api, seeded):
    response = api(FINANCE).get(
        "/api/finops/statement.pdf?project=other&period=2026-09"
    )

    assert response.status_code == 200


# ── month close ───────────────────────────────────────────────────────────────


def test_close_refuses_the_current_month(api):
    period = date.today().strftime("%Y-%m")

    assert api(FINANCE).post(f"/api/finops/close/{period}").status_code == 409


def test_close_refuses_a_month_with_non_final_invoice_lines(api, factory):
    with factory() as db:
        upsert(
            db,
            [_line(x_pricing_basis="invoice", x_is_final=False)],
        )

    assert api(FINANCE).post("/api/finops/close/2026-09").status_code == 409


def test_close_with_force_accepts_non_final_invoice_lines(api, factory):
    with factory() as db:
        upsert(db, [_line(x_pricing_basis="invoice", x_is_final=False)])

    response = api(FINANCE).post("/api/finops/close/2026-09?force=true")

    assert response.json()["closed_by"] == "fin"


def test_close_ignores_non_final_lines_that_are_not_invoices(api, factory):
    with factory() as db:
        upsert(db, [_line(x_is_final=False)])

    assert api(FINANCE).post("/api/finops/close/2026-09").status_code == 200


def test_close_twice_is_a_conflict(api, factory):
    with factory() as db:
        upsert(db, [_line()])
    api(FINANCE).post("/api/finops/close/2026-09")

    assert api(FINANCE).post("/api/finops/close/2026-09").status_code == 409


def test_closes_lists_the_closed_months(api, factory):
    with factory() as db:
        upsert(db, [_line()])
    api(ADMIN).post("/api/finops/close/2026-09")

    body = api(MEMBER).get("/api/finops/closes").json()

    assert [c["billing_period"] for c in body] == ["2026-09"]


def test_month_close_and_cost_centers_are_journaled():
    assert (
        match_action("POST", "/api/finops/close/2026-09")[0]
        == "finops.month.close"
    )
    assert match_action("PUT", "/api/finops/cost-centers")[0] == (
        "finops.cost_centers.update"
    )


# ── adjustments ───────────────────────────────────────────────────────────────


def _adjustments(db):
    return db.scalars(
        select(CostLineItem).where(
            CostLineItem.charge_category == "Adjustment"
        )
    ).all()


@pytest.fixture
def closed_sept(factory):
    with factory() as db:
        upsert(db, [_line(cost=1.0, start=datetime(2026, 9, 30))])
        db.add(MonthClose(billing_period="2026-09", closed_by="fin"))
        db.commit()
    return factory


def test_late_revision_of_a_closed_month_becomes_an_adjustment_on_the_next_one(
    closed_sept,
):
    with closed_sept() as db:
        upsert(db, [_line(cost=3.0, start=datetime(2026, 9, 30))])

        [adj] = _adjustments(db)

    assert (adj.x_cost_eur, adj.billing_period, adj.charge_period_start) == (
        2.0,
        "2026-10",
        datetime(2026, 10, 1),
    )
    assert adj.resource_id == "onprem/qcm/cpu#adj-2026-09-30"


def test_adjustment_keeps_the_closed_line_untouched(closed_sept):
    with closed_sept() as db:
        upsert(db, [_line(cost=3.0, start=datetime(2026, 9, 30))])

        original = db.scalar(
            select(CostLineItem).where(
                CostLineItem.billing_period == "2026-09"
            )
        )

    assert original.x_cost_eur == 1.0


def test_replaying_the_same_revision_does_not_stack_adjustments(closed_sept):
    revised = _line(cost=3.0, start=datetime(2026, 9, 30))
    with closed_sept() as db:
        upsert(db, [revised])
        upsert(db, [revised])
        upsert(db, [revised])

        [adj] = _adjustments(db)

    assert adj.x_cost_eur == 2.0


def test_a_second_revision_replaces_the_adjustment_delta(closed_sept):
    with closed_sept() as db:
        upsert(db, [_line(cost=3.0, start=datetime(2026, 9, 30))])
        upsert(db, [_line(cost=4.0, start=datetime(2026, 9, 30))])

        [adj] = _adjustments(db)

    assert adj.x_cost_eur == 3.0


def test_reverting_a_revision_zeroes_the_adjustment(closed_sept):
    with closed_sept() as db:
        upsert(db, [_line(cost=3.0, start=datetime(2026, 9, 30))])
        upsert(db, [_line(cost=1.0, start=datetime(2026, 9, 30))])

        [adj] = _adjustments(db)

    assert adj.x_cost_eur == 0.0


def test_an_unchanged_line_of_a_closed_month_books_nothing(closed_sept):
    with closed_sept() as db:
        written = upsert(db, [_line(cost=1.0, start=datetime(2026, 9, 30))])

        assert (written, _adjustments(db)) == (0, [])


def test_adjustment_lands_on_the_first_open_month(closed_sept):
    with closed_sept() as db:
        db.add(MonthClose(billing_period="2026-10", closed_by="fin"))
        db.commit()
        upsert(db, [_line(cost=2.0, start=datetime(2026, 9, 30))])

        [adj] = _adjustments(db)

    assert adj.billing_period == "2026-11"


def test_two_days_of_one_resource_get_distinct_adjustments(factory):
    days = [datetime(2026, 9, 29), datetime(2026, 9, 30)]
    with factory() as db:
        upsert(db, [_line(cost=1.0, start=d) for d in days])
        db.add(MonthClose(billing_period="2026-09", closed_by="fin"))
        db.commit()
        upsert(db, [_line(cost=2.0, start=d) for d in days])

        assert len(_adjustments(db)) == 2


# ── cost centers ──────────────────────────────────────────────────────────────


def test_finance_lists_every_project_in_cost_centers(api, seeded):
    body = api(FINANCE).get("/api/finops/cost-centers").json()

    assert {r["project_name"] for r in body} == {"qcm", "other"}


def test_member_only_sees_their_project_cost_center(api, seeded):
    body = api(MEMBER).get("/api/finops/cost-centers").json()

    assert body == [
        {"project_name": "qcm", "cost_center": "CC-1", "budget_owner": "Ana"}
    ]


def test_finance_updates_and_clears_cost_centers(api, seeded):
    client = api(FINANCE)

    client.put(
        "/api/finops/cost-centers",
        json={
            "centers": [
                {
                    "project_name": "other",
                    "cost_center": "CC-9",
                    "budget_owner": "Bo",
                },
                {"project_name": "qcm", "cost_center": ""},
            ]
        },
    )
    body = {
        r["project_name"]: r
        for r in client.get("/api/finops/cost-centers").json()
    }

    assert (body["other"]["cost_center"], body["qcm"]["cost_center"]) == (
        "CC-9",
        None,
    )


def test_cost_center_for_an_unknown_project_is_rejected(api, seeded):
    response = api(FINANCE).put(
        "/api/finops/cost-centers",
        json={"centers": [{"project_name": "ghost", "cost_center": "X"}]},
    )

    assert response.status_code == 400


# ── monthly report ────────────────────────────────────────────────────────────

PARIS = ZoneInfo("Europe/Paris")


@pytest.fixture
def report_db(seeded, monkeypatch):
    monkeypatch.setattr(monthly_report, "SessionLocal", seeded)
    monkeypatch.setattr(settings, "CMP_PUBLIC_URL", "https://cmp.test")
    return seeded


def test_report_message_is_french_and_links_the_finance_tab(report_db):
    with report_db() as db:
        message = monthly_report.format_report(
            monthly_report.build_report(db, "2026-09")
        )

    assert "Rapport FinOps 2026-09" in message
    assert "10,00 €" in message
    assert "4,00 € évités, 1,50 € réalisés" in message
    assert "https://cmp.test/finops/finance" in message


def test_report_lists_top_projects_and_unallocated_share(report_db):
    with report_db() as db:
        upsert(
            db,
            [_line("platform", 10.0, resource="plat/cpu")],
        )
        report = monthly_report.build_report(db, "2026-09")

    assert [p[0] for p in report["projects"]] == ["qcm", "other"]
    assert report["unallocated_pct"] == 50.0


def test_report_is_sent_once_for_the_previous_month(report_db):
    discord = mock.AsyncMock(return_value=True)
    now = datetime(2026, 10, 5, 9, 0, tzinfo=PARIS)

    with mock.patch.object(
        monthly_report.discord_service, "send_finops_monthly_report", discord
    ):
        first = asyncio.run(monthly_report.send_monthly_report(now))
        second = asyncio.run(monthly_report.send_monthly_report(now))

    assert (first, second, discord.await_count) == (True, False, 1)
    with report_db() as db:
        assert db.get(FinopsState, monthly_report.STATE_KEY).value == "2026-09"


def test_report_waits_for_the_fifth(report_db):
    discord = mock.AsyncMock(return_value=True)

    with mock.patch.object(
        monthly_report.discord_service, "send_finops_monthly_report", discord
    ):
        sent = asyncio.run(
            monthly_report.send_monthly_report(
                datetime(2026, 10, 4, 23, 0, tzinfo=PARIS)
            )
        )

    assert (sent, discord.await_count) == (False, 0)


def test_report_is_retried_when_discord_fails(report_db):
    discord = mock.AsyncMock(return_value=False)
    now = datetime(2026, 10, 5, 9, 0, tzinfo=PARIS)

    with mock.patch.object(
        monthly_report.discord_service, "send_finops_monthly_report", discord
    ):
        asyncio.run(monthly_report.send_monthly_report(now))

    with report_db() as db:
        assert db.get(FinopsState, monthly_report.STATE_KEY) is None


def test_report_follows_the_paris_date_not_utc(report_db):
    with report_db() as db:
        late_utc = datetime(2026, 10, 4, 23, 30, tzinfo=ZoneInfo("UTC"))

        assert monthly_report.is_due(db, late_utc) == "2026-09"


def test_csv_export_accepts_a_billing_month(api, seeded, monkeypatch):
    monkeypatch.setattr(settings, "FINOPS_PROVIDER", "real")

    response = api(FINANCE).get("/api/finops/export.csv?period=2026-09")

    assert response.text.count("\n") == 4
    assert "finops-all-2026-09.csv" in response.headers["content-disposition"]
