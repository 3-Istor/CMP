"""
Finance pack: monthly aggregates, XLSX / Parquet / PDF exports and month close.

Everything works on the FOCUS lines of one billing month (``billing_period``),
so an adjustment booked on a later month shows up in that later month.
"""

import io
import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timedelta

import pyarrow as pa
import pyarrow.parquet as pq
from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.finops import (
    CostCenter,
    CostLineItem,
    MonthClose,
    SavingsEntry,
)
from app.services.finops import platform
from app.services.finops.provider import resource_of

PERIOD_PATTERN = r"^\d{4}-(0[1-9]|1[0-2])$"
_PERIOD_RE = re.compile(PERIOD_PATTERN)
_NOT_A_PROJECT = ("platform", "__idle__")
_CAPEX_PREFIXES = ("amortissement", "hardware amortization")

BASIS_NOTE = (
    "Base tco (électricité on-prem) : montants TTC. "
    "Base invoice (factures cloud, SaaS) : montants HT. "
    "Base public_price : prix public catalogue."
)

# FOCUS column name, CostLineItem attribute, Parquet type.
FOCUS_COLUMNS = (
    ("ChargePeriodStart", "charge_period_start", pa.timestamp("us")),
    ("ChargePeriodEnd", "charge_period_end", pa.timestamp("us")),
    ("BillingPeriod", "billing_period", pa.string()),
    ("ProviderName", "provider_name", pa.string()),
    ("ServiceName", "service_name", pa.string()),
    ("ServiceCategory", "service_category", pa.string()),
    ("ChargeCategory", "charge_category", pa.string()),
    ("ResourceId", "resource_id", pa.string()),
    ("ResourceName", "resource_name", pa.string()),
    ("BilledCost", "billed_cost", pa.float64()),
    ("EffectiveCost", "effective_cost", pa.float64()),
    ("ListCost", "list_cost", pa.float64()),
    ("BillingCurrency", "billing_currency", pa.string()),
    ("ConsumedQuantity", "consumed_quantity", pa.float64()),
    ("ConsumedUnit", "consumed_unit", pa.string()),
    ("x_CostEur", "x_cost_eur", pa.float64()),
    ("x_FxRate", "x_fx_rate", pa.float64()),
    ("x_Project", "project", pa.string()),
    ("x_App", "app", pa.string()),
    ("x_Env", "env", pa.string()),
    ("x_Cloud", "cloud", pa.string()),
    ("x_AllocationMethod", "x_allocation_method", pa.string()),
    ("x_PricingBasis", "x_pricing_basis", pa.string()),
    ("x_EnergyKWh", "x_energy_kwh", pa.float64()),
    ("x_CO2eKg", "x_co2e_kg", pa.float64()),
    ("x_IsFinal", "x_is_final", pa.bool_()),
    ("x_IsEstimate", "x_is_estimate", pa.bool_()),
)


class CloseRefused(Exception):
    """The month cannot be closed; the message says why, in French."""


def valid_period(period: str) -> bool:
    return bool(_PERIOD_RE.match(period))


def current_period() -> str:
    return date.today().strftime("%Y-%m")


def previous_period(today: date) -> str:
    return (today.replace(day=1) - timedelta(days=1)).strftime("%Y-%m")


def nature(line: CostLineItem) -> str:
    """CapEx for hardware amortization, OpEx for everything else."""
    name = (line.service_name or "").lower()
    return "CapEx" if name.startswith(_CAPEX_PREFIXES) else "OpEx"


def is_unallocated(project: str | None) -> bool:
    return not project or project in _NOT_A_PROJECT


# ── Reads ─────────────────────────────────────────────────────────────────────


def month_lines(
    db: Session, period: str, projects: list[str] | None = None
) -> list[CostLineItem]:
    """Lines of *period*; ``projects`` restricts them, None means all."""
    query = select(CostLineItem).where(CostLineItem.billing_period == period)
    if projects is not None:
        query = query.where(CostLineItem.project.in_(projects))
    return list(
        db.scalars(
            query.order_by(CostLineItem.charge_period_start, CostLineItem.id)
        )
    )


def month_savings(
    db: Session, period: str, project: str | None = None
) -> list[SavingsEntry]:
    year, month = map(int, period.split("-"))
    first = date(year, month, 1)
    last = (first.replace(day=28) + timedelta(days=4)).replace(day=1)
    query = select(SavingsEntry).where(
        SavingsEntry.day >= first, SavingsEntry.day < last
    )
    if project:
        query = query.where(SavingsEntry.project == project)
    return list(
        db.scalars(query.order_by(SavingsEntry.day, SavingsEntry.project))
    )


def cost_centers(db: Session) -> dict[str, CostCenter]:
    return {c.project_name: c for c in db.scalars(select(CostCenter))}


def savings_totals(entries: list[SavingsEntry]) -> tuple[float, float]:
    """(avoided, realized): two figures that are never added together."""
    return (
        round(sum(e.avoided_eur or 0.0 for e in entries), 2),
        round(sum(e.realized_eur or 0.0 for e in entries), 2),
    )


# ── Aggregates ────────────────────────────────────────────────────────────────


def totals_by_provider(
    lines: list[CostLineItem],
) -> list[tuple[str, str, str, float]]:
    """(provider, pricing basis, CapEx/OpEx, cost) sorted by cost."""
    acc: dict[tuple[str, str, str], float] = defaultdict(float)
    for line in lines:
        key = (line.provider_name, line.x_pricing_basis, nature(line))
        acc[key] += line.x_cost_eur or 0.0
    return sorted(
        ((*k, round(v, 2)) for k, v in acc.items()), key=lambda r: -r[3]
    )


def totals_by_project(lines: list[CostLineItem]) -> list[tuple[str, float]]:
    acc: dict[str, float] = defaultdict(float)
    for line in lines:
        acc[line.project or "(non alloué)"] += line.x_cost_eur or 0.0
    return sorted(
        ((k, round(v, 2)) for k, v in acc.items()), key=lambda r: -r[1]
    )


def provider_sums(lines: list[CostLineItem]) -> list[tuple[str, float]]:
    acc: dict[str, float] = defaultdict(float)
    for line in lines:
        acc[line.provider_name] += line.x_cost_eur or 0.0
    return sorted(
        ((k, round(v, 2)) for k, v in acc.items()), key=lambda r: -r[1]
    )


def unallocated_pct(lines: list[CostLineItem]) -> float:
    total = sum(line.x_cost_eur or 0.0 for line in lines)
    if total <= 0:
        return 0.0
    loose = sum(
        line.x_cost_eur or 0.0
        for line in lines
        if is_unallocated(line.project)
    )
    return round(loose / total * 100, 1)


# ── XLSX ──────────────────────────────────────────────────────────────────────


def _safe(value):
    """A cell must never run as a formula when finance opens the file."""
    if isinstance(value, str) and value[:1] in (
        "=",
        "+",
        "-",
        "@",
        "\t",
        "\r",
    ):
        return "'" + value
    return value


def _append(sheet, values, bold: bool = False) -> None:
    sheet.append([_safe(v) for v in values])
    if bold:
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)


def _widths(sheet, widths: tuple[int, ...]) -> None:
    for index, width in enumerate(widths):
        sheet.column_dimensions[chr(ord("A") + index)].width = width


def cluster_days(lines: list[CostLineItem]) -> dict[date, platform.ClusterDay]:
    return platform.cluster_days(
        (
            line.charge_period_start.date(),
            line.resource_name,
            line.project,
            line.x_cost_eur or 0.0,
        )
        for line in lines
        if line.x_source == "opencost" and line.cloud == "onprem"
    )


def build_xlsx(
    period: str,
    lines: list[CostLineItem],
    savings: list[SavingsEntry],
    centers: dict[str, CostCenter],
    gaps: dict[str, dict],
) -> bytes:
    """``gaps`` is ``reconciliation.reconcile`` per cloud provider."""
    book = Workbook()
    cluster = cluster_days(lines)
    fees = platform.fees_by_project(cluster)

    summary = book.active
    summary.title = "Synthèse"
    _append(summary, [f"Showback CNP, période {period}"], bold=True)
    _append(summary, [BASIS_NOTE])
    _append(summary, [])
    _append(summary, ["Par fournisseur"], bold=True)
    _append(
        summary, ["Fournisseur", "Base", "Nature", "Coût (EUR)"], bold=True
    )
    for provider, basis, kind, cost in totals_by_provider(lines):
        _append(summary, [provider, basis, kind, cost])
    _append(summary, [])
    _append(summary, ["Par projet"], bold=True)
    _append(
        summary, ["Projet", "Centre de coûts", "Base", "Coût (EUR)"], bold=True
    )
    bases: dict[str, set[str]] = defaultdict(set)
    for line in lines:
        bases[line.project or "(non alloué)"].add(line.x_pricing_basis)
    for project, cost in totals_by_project(lines):
        center = centers.get(project)
        _append(
            summary,
            [
                project,
                center.cost_center if center else "",
                ", ".join(sorted(bases[project])),
                cost,
            ],
        )
    _widths(summary, (28, 18, 18, 16))

    projects = book.create_sheet("Projets")
    _append(
        projects,
        [
            "Projet",
            "Centre de coûts",
            "Responsable budget",
            "Fournisseur",
            "Base",
            "Coût (EUR)",
            "Dont forfait plateforme (EUR)",
        ],
        bold=True,
    )
    grid: dict[tuple[str, str, str], float] = defaultdict(float)
    for line in lines:
        key = (
            line.project or "(non alloué)",
            line.provider_name,
            line.x_pricing_basis,
        )
        grid[key] += line.x_cost_eur or 0.0
    fee_shown: set[str] = set()
    for (project, provider, basis), cost in sorted(grid.items()):
        center = centers.get(project)
        fee = 0.0 if project in fee_shown else fees.get(project, 0.0)
        fee_shown.add(project)
        _append(
            projects,
            [
                project,
                center.cost_center if center else "",
                (center.budget_owner or "") if center else "",
                provider,
                basis,
                round(cost, 2),
                fee,
            ],
        )
    _widths(projects, (28, 18, 24, 28, 16, 14, 18))

    shared = book.create_sheet("Plateforme")
    first_day = min(cluster, default=date.today())
    last_day = max(cluster, default=first_day)
    summary_report = platform.report(cluster, first_day, last_day)
    _append(shared, ["Coût de la plateforme"], bold=True)
    _append(
        shared,
        [
            "Total plateforme et capacité inoccupée (EUR)",
            summary_report["total_eur"],
        ],
    )
    _append(
        shared, ["Dont capacité inoccupée (EUR)", summary_report["idle_eur"]]
    )
    _append(
        shared,
        ["Recouvert par les forfaits (EUR)", summary_report["recovered_eur"]],
    )
    _append(shared, ["Non couvert (EUR)", summary_report["uncovered_eur"]])
    _append(
        shared, ["Emplacements d'applications", summary_report["app_slots"]]
    )
    _append(
        shared,
        ["Applications actives (moyenne)", summary_report["active_apps_avg"]],
    )
    _append(
        shared,
        [
            "Forfait par application et par mois (EUR)",
            summary_report["fee_per_app_month_eur"],
        ],
    )
    _append(shared, [])
    _append(shared, ["Namespace", "Coût (EUR)"], bold=True)
    for entry in summary_report["by_namespace"]:
        _append(shared, [entry["namespace"], entry["eur"]])
    _append(shared, [])
    _append(
        shared,
        [
            "Jour",
            "Plateforme (EUR)",
            "Inoccupé (EUR)",
            "Recouvert (EUR)",
            "Non couvert (EUR)",
        ],
        bold=True,
    )
    for day in summary_report["daily"]:
        _append(
            shared,
            [
                day["date"],
                day["platform_eur"],
                day["idle_eur"],
                day["recovered_eur"],
                day["uncovered_eur"],
            ],
        )
    _widths(shared, (44, 18, 16, 16, 18))

    providers = book.create_sheet("Fournisseurs")
    _append(
        providers,
        ["Fournisseur", "Catégorie", "Base", "Nature", "Lignes", "Coût (EUR)"],
        bold=True,
    )
    by_service: dict[tuple[str, str, str, str], list[float]] = defaultdict(
        lambda: [0, 0.0]
    )
    for line in lines:
        key = (
            line.provider_name,
            line.service_category,
            line.x_pricing_basis,
            nature(line),
        )
        by_service[key][0] += 1
        by_service[key][1] += line.x_cost_eur or 0.0
    for key, (count, cost) in sorted(by_service.items()):
        _append(providers, [*key, count, round(cost, 2)])
    _widths(providers, (28, 24, 16, 10, 10, 14))

    reconciled = book.create_sheet("Écarts")
    _append(
        reconciled,
        ["Facture du cloud comparée à sa répartition (EUR)."],
    )
    _append(
        reconciled,
        [
            "Fournisseur",
            "Facture",
            "Alloué aux projets",
            "Inactif",
            "Plateforme",
            "Non alloué",
            "Non alloué (%)",
            "Écart inexpliqué",
        ],
        bold=True,
    )
    for provider, gap in sorted(gaps.items()):
        _append(
            reconciled,
            [
                provider,
                round(gap["invoice_total_eur"], 2),
                round(gap["allocated_eur"], 2),
                round(gap["idle_eur"], 2),
                round(gap["platform_eur"], 2),
                round(gap["unallocated_eur"], 2),
                gap["unallocated_pct"],
                round(gap["gap_eur"], 2),
            ],
        )
    _widths(reconciled, (12, 12, 18, 10, 12, 12, 14, 16))

    ledger = book.create_sheet("Économies")
    _append(
        ledger,
        [
            "Jour",
            "Projet",
            "Application",
            "Catégorie",
            "Évité (EUR)",
            "Réalisé (EUR)",
            "kWh évités",
            "kgCO2e évités",
        ],
        bold=True,
    )
    for entry in savings:
        _append(
            ledger,
            [
                entry.day,
                entry.project,
                entry.app or "",
                entry.category,
                entry.avoided_eur or 0.0,
                entry.realized_eur or 0.0,
                entry.energy_avoided_kwh or 0.0,
                entry.co2e_avoided_kg or 0.0,
            ],
        )
    avoided, realized = savings_totals(savings)
    _append(
        ledger, ["Total", "", "", "", avoided, realized, "", ""], bold=True
    )
    _append(
        ledger,
        ["Le coût évité et le coût réalisé ne s'additionnent jamais."],
    )
    _widths(ledger, (12, 24, 24, 24, 14, 14, 12, 14))

    detail = book.create_sheet("Lignes")
    _append(
        detail,
        [name for name, _attr, _type in FOCUS_COLUMNS] + ["Nature"],
        bold=True,
    )
    for line in lines:
        _append(
            detail,
            [getattr(line, attr) for _n, attr, _t in FOCUS_COLUMNS]
            + [nature(line)],
        )

    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


# ── Parquet ───────────────────────────────────────────────────────────────────


def build_parquet(lines: list[CostLineItem]) -> bytes:
    schema = pa.schema(
        [pa.field(name, kind) for name, _attr, kind in FOCUS_COLUMNS]
    )
    table = pa.table(
        {
            name: [getattr(line, attr) for line in lines]
            for name, attr, _kind in FOCUS_COLUMNS
        },
        schema=schema,
    )
    out = io.BytesIO()
    pq.write_table(table, out)
    return out.getvalue()


# ── PDF ───────────────────────────────────────────────────────────────────────


def _latin(text: str) -> str:
    # Core PDF fonts are Latin-1: no euro sign, no typographic quotes.
    return (
        text.replace("€", "EUR")
        .replace("’", "'")
        .encode("latin-1", "replace")
        .decode("latin-1")
    )


@dataclass
class _Table:
    title: str
    header: tuple[str, ...]
    rows: list[tuple]
    widths: tuple[int, ...]


def _draw(pdf: FPDF, table: _Table) -> None:
    pdf.set_font("Helvetica", "B", 10)
    pdf.cell(0, 6, _latin(table.title), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "B", 8)
    for width, name in zip(table.widths, table.header):
        pdf.cell(width, 5, _latin(name), border="B")
    pdf.ln()
    pdf.set_font("Helvetica", "", 8)
    if not table.rows:
        pdf.cell(0, 5, "Aucune donnee.", new_x="LMARGIN", new_y="NEXT")
    for row in table.rows:
        for width, value in zip(table.widths, row):
            text = f"{value:.2f}" if isinstance(value, float) else str(value)
            pdf.cell(width, 5, _latin(text)[:48], align="")
        pdf.ln()
    pdf.ln(2)


def build_statement_pdf(
    project: str,
    period: str,
    lines: list[CostLineItem],
    savings: list[SavingsEntry],
    center: CostCenter | None,
    budget_eur: float | None,
) -> bytes:
    """One-page showback statement of a project for a billing month."""
    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(False)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 15)
    pdf.cell(
        0,
        9,
        _latin(f"Releve showback {project}"),
        new_x="LMARGIN",
        new_y="NEXT",
    )
    pdf.set_font("Helvetica", "", 9)
    total = round(sum(line.x_cost_eur or 0.0 for line in lines), 2)
    pdf.cell(0, 5, f"Periode : {period}", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(
        0,
        5,
        _latin(
            "Centre de couts : "
            + (center.cost_center if center else "non renseigne")
            + " | Responsable budget : "
            + (
                (center.budget_owner or "non renseigne")
                if center
                else "non renseigne"
            )
        ),
        new_x="LMARGIN",
        new_y="NEXT",
    )
    budget_text = (
        f" (budget mensuel {budget_eur:.2f} EUR)" if budget_eur else ""
    )
    pdf.set_font("Helvetica", "B", 11)
    pdf.cell(
        0,
        8,
        _latin(f"Total : {total:.2f} EUR{budget_text}"),
        new_x="LMARGIN",
        new_y="NEXT",
    )

    by_resource: dict[str, float] = defaultdict(float)
    by_app: dict[str, float] = defaultdict(float)
    for line in lines:
        by_resource[resource_of(line.service_name, line.service_category)] += (
            line.x_cost_eur or 0.0
        )
        by_app[line.app or "(projet)"] += line.x_cost_eur or 0.0
    avoided, realized = savings_totals(savings)

    for table in (
        _Table(
            "Par ressource",
            ("Ressource", "Cout (EUR)"),
            [(k, round(v, 2)) for k, v in sorted(by_resource.items())],
            (80, 40),
        ),
        _Table(
            "Par fournisseur",
            ("Fournisseur", "Base", "Cout (EUR)"),
            [(p, b, c) for p, b, _n, c in totals_by_provider(lines)],
            (80, 40, 40),
        ),
        _Table(
            "Principales applications",
            ("Application", "Cout (EUR)"),
            sorted(
                ((k, round(v, 2)) for k, v in by_app.items()),
                key=lambda r: -r[1],
            )[:5],
            (80, 40),
        ),
        _Table(
            "Economies du mois (jamais additionnees)",
            ("Cout evite (EUR)", "Cout realise (EUR)"),
            [(avoided, realized)],
            (80, 80),
        ),
    ):
        _draw(pdf, table)

    pdf.set_font("Helvetica", "", 7)
    pdf.multi_cell(
        0,
        3.5,
        _latin(
            "Methode : les couts on-prem sont repartis par OpenCost a partir "
            "d'un TCO (electricite TTC, materiel gratuit). Les factures "
            "cloud sont HT. Cout evite = cout alloue libere (pods "
            "endormis) ; cout realise = ce qui a reellement quitte la "
            "facture. "
            + (
                "Certaines lignes reposent sur des prix "
                "provisoires (estimation)."
                if any(line.x_is_estimate for line in lines)
                else ""
            )
        ),
    )
    return bytes(pdf.output())


# ── Month close ───────────────────────────────────────────────────────────────


def closed_months(db: Session) -> list[MonthClose]:
    return list(
        db.scalars(
            select(MonthClose).order_by(MonthClose.billing_period.desc())
        )
    )


def close_month(
    db: Session, period: str, user: str, force: bool = False
) -> MonthClose:
    if period >= current_period():
        raise CloseRefused(
            "Le mois en cours ou à venir ne peut pas être clôturé."
        )
    if db.scalar(
        select(MonthClose).where(MonthClose.billing_period == period)
    ):
        raise CloseRefused(f"Le mois {period} est déjà clôturé.")
    lines = month_lines(db, period)
    pending = [
        line
        for line in lines
        if line.x_pricing_basis == "invoice" and not line.x_is_final
    ]
    if not force and not lines:
        raise CloseRefused(
            f"Aucune ligne de coût pour {period}. Forcez pour clôturer quand même."
        )
    if not force and pending:
        raise CloseRefused(
            f"{len(pending)} ligne(s) de facture non définitive(s) en {period}. "
            "Forcez pour clôturer quand même."
        )
    closed = MonthClose(billing_period=period, closed_by=user)
    db.add(closed)
    db.commit()
    db.refresh(closed)
    return closed
